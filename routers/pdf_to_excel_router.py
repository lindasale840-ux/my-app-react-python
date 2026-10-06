import io
import os
import zipfile
import re
from typing import List, Optional
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse, JSONResponse
import fitz  # PyMuPDF
import pdfplumber
import cv2
import numpy as np
from PIL import Image
import pytesseract
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from pydantic import BaseModel

router = APIRouter(prefix="/api/pdf-to-excel", tags=["PdfToExcelExport"])

# Helper sanitize sheet name
def sanitize_sheet_title(name: str) -> str:
    cleaned = re.sub(r'[\\/*?:\[\]]', '_', name)
    return cleaned[:30] if cleaned else "Sheet1"

# Pydantic model cho matrix từ Frontend
class CellModel(BaseModel):
    row: int
    col: int
    rowspan: int = 1
    colspan: int = 1
    text: str = ""

class MatrixPayload(BaseModel):
    filename: str
    rows: int
    cols: int
    cells: List[CellModel]

class BatchExportRequest(BaseModel):
    export_mode: str  # "single_file" hoặc "zip_archive"
    matrices: List[MatrixPayload]


def extract_table_matrix_from_pdf(pdf_bytes: bytes, filename: str):
    """
    Bản sửa lỗi: Duyệt TOÀN BỘ các trang của file PDF thay vì chỉ đọc trang đầu tiên (pages[0])
    """
    bg_image_b64 = ""
    cells_list = []
    current_row_offset = 0
    max_cols = 1

    try:
        # 1. Render ảnh nền trang đầu làm đại diện (hoặc ghép ảnh nếu cần)
        pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        total_pages = len(pdf_doc)
        
        if total_pages > 0:
            page_0 = pdf_doc[0]
            pix = page_0.get_pixmap(dpi=100)
            img_bytes = pix.tobytes("png")
            import base64
            bg_image_b64 = "data:image/png;base64," + base64.b64encode(img_bytes).decode('utf-8')

        # 2. VÒNG LẶP DUYỆT TẤT CẢ CÁC TRANG CỦA PDF
        try:
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as plumber_pdf:
                for page_idx, page in enumerate(plumber_pdf.pages):
                    tables = page.extract_tables()
                    
                    if tables:
                        for table in tables:
                            if not table:
                                continue
                            
                            # Cập nhật số cột lớn nhất
                            for r in table:
                                if len(r) > max_cols:
                                    max_cols = len(r)

                            # Đưa dữ liệu trang này vào danh sách chung (tăng nối tiếp số hàng)
                            for r_idx, row in enumerate(table):
                                actual_row = current_row_offset + r_idx + 1
                                for c_idx, val in enumerate(row):
                                    text_val = str(val).strip() if val is not None else ""
                                    cells_list.append({
                                        "row": actual_row,
                                        "col": c_idx + 1,
                                        "rowspan": 1,
                                        "colspan": 1,
                                        "text": text_val
                                    })
                            
                            # Tăng offset hàng cho trang/bảng tiếp theo
                            current_row_offset += len(table)

        except Exception as e_plumber:
            print(f"pdfplumber error, fallbacking: {e_plumber}")

        # 3. Fallback nếu không bóc tách được bảng nào
        if not cells_list:
            current_row_offset = 5
            max_cols = 4
            for r in range(1, current_row_offset + 1):
                for c in range(1, max_cols + 1):
                    cells_list.append({
                        "row": r,
                        "col": c,
                        "rowspan": 1,
                        "colspan": 1,
                        "text": f"Ô ({r},{c})"
                    })

        return {
            "filename": filename,
            "rows": current_row_offset if current_row_offset > 0 else 5,
            "cols": max_cols,
            "cells": cells_list,
            "bg_image": bg_image_b64,
            "total_pages": total_pages, # Trả thêm tổng số trang để Frontend biết
            "status": "Ready"
        }

    except Exception as e:
        print(f"Lỗi extract_table_matrix_from_pdf: {e}")
        return {
            "filename": filename,
            "rows": 3,
            "cols": 3,
            "cells": [
                {"row": 1, "col": 1, "rowspan": 1, "colspan": 1, "text": "Lỗi đọc bảng"},
                {"row": 1, "col": 2, "rowspan": 1, "colspan": 2, "text": str(e)}
            ],
            "bg_image": "",
            "status": "Error",
            "error_msg": str(e)
        }


@router.post("/analyze-batch")
async def api_analyze_batch(files: List[UploadFile] = File(...)):
    """
    Phân tích danh sách PDF tải lên và trả về ma trận cấu trúc + ảnh xem trước
    """
    results = []
    for file in files:
        content = await file.read()
        res = extract_table_matrix_from_pdf(content, file.filename)
        results.append(res)
    return JSONResponse(content={"status": "success", "data": results})


# Cập nhật hàm write_matrix_to_sheet trong pdf_to_excel_router.py
def write_matrix_to_sheet(ws, matrix: MatrixPayload):
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    header_fill = PatternFill(start_color="007BFF", end_color="007BFF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    merged_set = set()
    
    # Ghi ma trận dữ liệu
    for cell in matrix.cells:
        r, c = cell.row, cell.col
        r_span, c_span = cell.rowspan, cell.colspan
        
        if (r, c) in merged_set:
            continue
            
        cell_obj = ws.cell(row=r, column=c)
        val = cell.text.strip() if cell.text else ""
        
        # Parse số / định dạng
        if val.replace('.', '', 1).replace(',', '', 1).isdigit():
            try:
                cell_obj.value = float(val.replace(',', '.')) if ',' in val else float(val)
            except ValueError:
                cell_obj.value = val
        else:
            cell_obj.value = val
            
        cell_obj.alignment = Alignment(vertical='center', wrap_text=True)
        cell_obj.border = thin_border
        
        # Nếu là hàng tiêu đề (Hàng 1)
        if r == 1:
            cell_obj.fill = header_fill
            cell_obj.font = header_font
            
        if r_span > 1 or c_span > 1:
            end_r = r + r_span - 1
            end_c = c + c_span - 1
            ws.merge_cells(start_row=r, start_column=c, end_row=end_r, end_column=end_c)
            for mr in range(r, end_r + 1):
                for mc in range(c, end_c + 1):
                    merged_set.add((mr, mc))
                    ws.cell(row=mr, column=mc).border = thin_border


@router.post("/export-excel")
async def api_export_excel(payload: BatchExportRequest):
    """
    Xuất Excel từ ma trận JSON kiểm duyệt từ Frontend (Single Multi-Sheet HOẶC ZIP File)
    """
    try:
        if payload.export_mode == "single_file":
            wb = openpyxl.Workbook()
            wb.remove(wb.active)  # Xóa sheet mặc định
            
            for idx, mat in enumerate(payload.matrices):
                sheet_title = sanitize_sheet_title(os.path.splitext(mat.filename)[0])
                ws = wb.create_sheet(title=sheet_title)
                write_matrix_to_sheet(ws, mat)
                
            stream = io.BytesIO()
            wb.save(stream)
            stream.seek(0)
            
            return StreamingResponse(
                stream,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=Exported_PDF_Data.xlsx"}
            )
            
        else:  # zip_archive
            zip_stream = io.BytesIO()
            with zipfile.ZipFile(zip_stream, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for mat in payload.matrices:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Sheet1"
                    write_matrix_to_sheet(ws, mat)
                    
                    excel_stream = io.BytesIO()
                    wb.save(excel_stream)
                    excel_stream.seek(0)
                    
                    base_name = os.path.splitext(mat.filename)[0]
                    zip_file.writestr(f"{base_name}.xlsx", excel_stream.getvalue())
                    
            zip_stream.seek(0)
            return StreamingResponse(
                zip_stream,
                media_type="application/zip",
                headers={"Content-Disposition": "attachment; filename=Exported_Excel_Files.zip"}
            )
            
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi xuất Excel: {str(e)}")